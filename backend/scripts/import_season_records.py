"""赛事纪录 Excel 导入脚本（幂等，可重复执行）。

把「赛事纪录汇总.xlsx」中的真实比赛导入数据库：
- 球队：按主/客队名称去重自动建队（分组为「学院联赛」）
- 比赛：每行一场，历史比赛一律置为已结束，时间补 15:00，比分取自「比分」列
- 进球人/红黄牌列是叙事文本，无法可靠拆分到球员，故不解析，可在系统内人工补录

用法：
  cd backend
  python scripts/import_season_records.py "C:\\path\\赛事纪录汇总.xlsx"
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from app.database import SessionLocal, engine  # noqa: E402
from app.models import Base, Match, Team  # noqa: E402
from app.seed import seed_db  # noqa: E402

GROUP_NAME = "学院联赛"
DEFAULT_TIME = "15:00"

# 球队别名归一化：同一学院在纪录中的多种写法 → 规范名。
# 只合并写法差异明显的变体；不同年级/性质的队伍（如体教21/22级队）保留独立。
TEAM_ALIASES = {
    "AI学院": "人工智能学院",
    "人工智能(AI)学院": "人工智能学院",
    "人工智能（AI）学院": "人工智能学院",
    "农学与生物科学学院": "农学与生物科技学院",
    "农学与生物科技院": "农学与生物科技学院",
    "创新创业含弘学院": "创新创业学院",
    "创新创业学院 含弘学院": "创新创业学院",
    "创新创业学院含弘学院": "创新创业学院",
    "动物医学院": "动物医学学院",
    "动物科学与技术学院": "动物科学学院",
    "动物科学技术学院": "动物科学学院",
    "动物科技学院": "动物科学学院",
    "动科": "动物科学学院",
    "历史与文化学院": "历史文化学院",
    "历史与文化学院 民族学院": "历史文化学院",
    "历史文化学院 民族学院": "历史文化学院",
    "历史文化学院民族学院": "历史文化学院",
    "商贸院": "商贸学院",
    "园林园艺学院": "园艺园林学院",
    "地理与科学学院": "地理科学学院",
    "地理科学与技术学院": "地理科学学院",
    "心理学部": "心理学院",
    "数学与统计学学院": "数学与统计学院",
    "数学统计学院": "数学与统计学院",
    "新闻与传媒学院": "新闻传媒学院",
    "材料能源学院": "材料与能源学院",
    "桑蚕纺织与生物质学院": "蚕桑纺织与生物质科学学院",
    "桑蚕纺织与生物质科学院": "蚕桑纺织与生物质科学学院",
    "蚕桑纺织与生物质学院": "蚕桑纺织与生物质科学学院",
    "蚕桑纺织与生物质科学学院院": "蚕桑纺织与生物质科学学院",
    "蚕桑纺织生物质科学学院": "蚕桑纺织与生物质科学学院",
    "蚕生院": "蚕桑纺织与生物质科学学院",
    "物理学院": "物理科学与技术学院",
    "物理科学与技术学院院": "物理科学与技术学院",
    "物理科学学院": "物理科学与技术学院",
    "物理科学技术学院": "物理科学与技术学院",
    "电信院": "电子信息工程学院",
    "电子信息学院": "电子信息工程学院",
    "电子信息工程技术学院": "电子信息工程学院",
    "经济管理学院学院": "经济管理学院",
    "西塔足球队": "西塔学院",
    "计信院": "计算机与信息科学学院",
    "计算机与信息工程学院": "计算机与信息科学学院",
    "计算机与信息技术学院": "计算机与信息科学学院",
    "计算机与信息技术学院 软件学院": "计算机与信息科学学院",
    "计算机与信息科学学院 软件学院": "计算机与信息科学学院",
    "计算机信息科学学院": "计算机与信息科学学院",
    "计算机信息科学学院 软件学院": "计算机与信息科学学院",
    "资源与环境学院": "资源环境学院",
    "植保学院": "植物保护学院",
    "生物科学学院": "生命科学学院",
    "运动训练2022": "体育学院2022级运动训练专业",
    "体院研究生": "体育学院研究生队",
    "体院研究生队": "体育学院研究生队",
    "体育学院研究生": "体育学院研究生队",
    "体育学院研究生联队": "体育学院研究生队",
    "体育研究生": "体育学院研究生队",
}

COL_INDEX = {
    "seq": 0, "source": 1, "date": 2, "venue": 3,
    "home": 4, "away": 5, "score": 6,
}


def import_records(xlsx_path: str) -> None:
    from openpyxl import load_workbook

    Base.metadata.create_all(bind=engine)  # 确保表结构存在（幂等）
    db = SessionLocal()
    try:
        seed_db(db)  # 演示种子数据
        existing_teams = {team.name: team for team in db.query(Team).all()}
        created_teams: list[str] = []
        imported, skipped = 0, []

        workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        for row_index, row in enumerate(rows[1:], start=2):  # 跳过表头
            if row[COL_INDEX["home"]] is None or row[COL_INDEX["away"]] is None:
                continue
            home_name = _canonical(str(row[COL_INDEX["home"]]))
            away_name = _canonical(str(row[COL_INDEX["away"]]))
            date_text = str(row[COL_INDEX["date"]] or "").strip()
            score_text = str(row[COL_INDEX["score"]] or "").strip()
            if not home_name or not away_name or not date_text:
                skipped.append(f"第{row_index}行：主/客队或日期缺失")
                continue

            score = _parse_score(score_text)
            if score is None:
                skipped.append(f"第{row_index}行：比分无法解析「{score_text}」")
                continue

            home_id = _ensure_team(db, existing_teams, created_teams, home_name)
            away_id = _ensure_team(db, existing_teams, created_teams, away_name)
            year = date_text[:4]
            match_id = f"m{year}{int(row[COL_INDEX['seq']]):04d}"
            if db.get(Match, match_id) is not None:
                continue  # 已导入，跳过（幂等）

            db.add(Match(
                id=match_id, stage="group", group_name=GROUP_NAME,
                home_team_id=home_id, away_team_id=away_id,
                home_score=score[0], away_score=score[1],
                time=f"{date_text} {DEFAULT_TIME}",
                venue=str(row[COL_INDEX["venue"]] or "未提及").strip(),
                status="finished",
            ))
            imported += 1

        db.commit()
        print(f"导入完成：新建球队 {len(created_teams)} 支，导入比赛 {imported} 场")
        for item in skipped[:20]:
            print(f"  - 跳过：{item}")
        if len(skipped) > 20:
            print(f"  ... 其余 {len(skipped) - 20} 条跳过原因略")
    finally:
        db.close()


def _canonical(name: str) -> str:
    """归一化球队名：去空格 + 别名映射。"""
    cleaned = " ".join(name.split())
    return TEAM_ALIASES.get(cleaned, cleaned)


def _ensure_team(db, existing_teams: dict, created_teams: list, name: str) -> str:
    """按名称取球队ID；不存在则创建并记录。"""
    if name in existing_teams:
        return existing_teams[name].id
    team = Team(id=_team_id(name), name=name, group=GROUP_NAME, captain_id=None)
    db.add(team)
    existing_teams[name] = team
    created_teams.append(name)
    return team.id


def _team_id(name: str) -> str:
    import hashlib

    digest = hashlib.md5(name.encode("utf-8")).hexdigest()[:8]
    return f"t_{digest}"


def _parse_score(text: str) -> tuple[int, int] | None:
    match = re.search(r"(\d{1,2})\s*[:：]\s*(\d{1,2})", text)
    if match is None:
        return None
    return int(match.group(1)), int(match.group(2))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法：python scripts/import_season_records.py <xlsx路径>")
        sys.exit(1)
    import_records(sys.argv[1])
